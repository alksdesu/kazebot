from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from .storage import MaterialError, bounded_text


def _title(spec: dict) -> str:
    return bounded_text(spec.get("title", "未命名文档"), "标题", 160, required=True)


def _sections(spec: dict) -> list[dict]:
    sections = spec.get("sections", [])
    if not isinstance(sections, list) or len(sections) > 100:
        raise MaterialError("sections 必须为最多 100 个章节的列表。")
    result = []
    total = 0
    for section in sections:
        if not isinstance(section, dict):
            raise MaterialError("每个章节须为对象。")
        heading = bounded_text(section.get("heading", ""), "章节标题", 300)
        paragraphs = section.get("paragraphs", [])
        if not isinstance(paragraphs, list):
            raise MaterialError("paragraphs 必须为列表。")
        paragraphs = [bounded_text(text, "段落", 10000) for text in paragraphs]
        total += sum(map(len, paragraphs))
        if total > 100000:
            raise MaterialError("正文超过 10 万字。", "limit_exceeded", 413)
        table = section.get("table", [])
        if not isinstance(table, list) or len(table) > 500:
            raise MaterialError("表格超过 500 行或格式错误。")
        rows = []
        for row in table:
            if not isinstance(row, list) or len(row) > 20:
                raise MaterialError("表格每行须最多 20 列。")
            rows.append([bounded_text(str(cell), "单元格", 1000) for cell in row])
        result.append({"heading": heading, "paragraphs": paragraphs, "table": rows})
    if not result:
        raise MaterialError("请提供至少一个章节。")
    return result


def _pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    configured = os.environ.get("CLONOTH_MATERIALS_FONT", "").strip()
    if configured:
        from reportlab.pdfbase.ttfonts import TTFont
        pdfmetrics.registerFont(TTFont("Materials", configured))
        return "Materials"
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def build_file(kind: str, spec: dict, directory: Path) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    title = _title(spec)
    path = directory / f"document.{kind}"
    if kind == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

        font = _pdf_font()
        normal = ParagraphStyle("body", fontName=font, fontSize=11, leading=17, spaceAfter=8)
        heading = ParagraphStyle("heading", parent=normal, fontSize=15, leading=22, spaceBefore=12)
        elements = [Paragraph(escape(title), ParagraphStyle("title", parent=normal, fontSize=22, leading=28)), Spacer(1, 8 * mm)]
        for section in _sections(spec):
            if section["heading"]:
                elements.append(Paragraph(escape(section["heading"]), heading))
            elements.extend(Paragraph(escape(p).replace("\n", "<br/>"), normal) for p in section["paragraphs"])
            if section["table"]:
                rows = [[Paragraph(escape(cell), normal) for cell in row] for row in section["table"]]
                table = Table(rows, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.lightgrey),
                                           ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                elements.append(table)
        SimpleDocTemplate(str(path), title=title, leftMargin=20 * mm, rightMargin=20 * mm).build(elements)
    elif kind == "docx":
        from docx import Document
        from docx.shared import Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(11)
        normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        document.add_heading(title, 0)
        for section in _sections(spec):
            if section["heading"]:
                document.add_heading(section["heading"], 1)
            for text in section["paragraphs"]:
                document.add_paragraph(text)
            if section["table"]:
                width = max(map(len, section["table"]))
                if not width:
                    continue
                table = document.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for row in section["table"]:
                    cells = table.add_row().cells
                    for index, text in enumerate(row):
                        cells[index].text = text
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
        document.save(path)
        Document(path)
    elif kind == "xlsx":
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        sheets = spec.get("sheets", [])
        if not isinstance(sheets, list) or not 1 <= len(sheets) <= 10:
            raise MaterialError("请提供 1–10 张工作表。")
        workbook = Workbook()
        workbook.remove(workbook.active)
        count = 0
        for index, data in enumerate(sheets, 1):
            name = bounded_text(data.get("name", f"Sheet{index}"), "工作表名", 31, required=True)
            if any(char in name for char in "[]:*?/\\") or name in workbook.sheetnames:
                raise MaterialError("工作表名不合法或重复。")
            sheet = workbook.create_sheet(name)
            rows = data.get("rows", [])
            if not isinstance(rows, list) or len(rows) > 5000:
                raise MaterialError("工作表最多 5000 行。")
            for row_i, row in enumerate(rows, 1):
                if not isinstance(row, list) or len(row) > 100:
                    raise MaterialError("工作表每行最多 100 列。")
                for col_i, item in enumerate(row, 1):
                    count += 1
                    if count > 50000:
                        raise MaterialError("工作簿超过 50000 个单元格。", "limit_exceeded", 413)
                    cell = sheet.cell(row_i, col_i)
                    if isinstance(item, dict) and "formula" in item:
                        formula = bounded_text(item["formula"], "公式", 500)
                        functions = re.findall(r"([A-Za-z_][A-Za-z_0-9.]*)\s*\(", formula)
                        if not formula.startswith("=") or "[" in formula or any(f.upper() not in {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "IF", "ROUND", "ABS"} for f in functions):
                            raise MaterialError("公式含不支持的函数或外部引用。")
                        cell.value = formula
                    elif item is None or isinstance(item, (int, float, bool)):
                        cell.value = item
                    elif isinstance(item, str):
                        cell.value = bounded_text(item, "单元格", 10000)
                        cell.data_type = "s"
                    else:
                        raise MaterialError("单元格须为文字、数值或显式公式对象。")
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if row_i == 1:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="E7EDF3")
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(40, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = "1:1"
        workbook.save(path)
        load_workbook(path).close()
    elif kind == "pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt

        slides = spec.get("slides", [])
        if not isinstance(slides, list) or not 1 <= len(slides) <= 30:
            raise MaterialError("请提供 1–30 页幻灯片。")
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        for data in slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = bounded_text(data.get("title", ""), "幻灯片标题", 120, required=True)
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
            bullets = data.get("bullets", [])
            if not isinstance(bullets, list) or len(bullets) > 8:
                raise MaterialError("每页最多 8 条正文。")
            frame = slide.placeholders[1].text_frame
            frame.clear()
            frame.word_wrap = True
            for index, item in enumerate(bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = bounded_text(item, "幻灯片正文", 200)
                paragraph.font.size = Pt(22)
            notes = data.get("notes", "")
            if notes:
                slide.notes_slide.notes_text_frame.text = bounded_text(notes, "讲者备注", 10000)
        presentation.core_properties.title = title
        presentation.save(path)
        Presentation(path)
    else:
        raise MaterialError("生成格式须为 pdf/docx/xlsx/pptx。", "unsupported_type", 415)
    return path
