from __future__ import annotations

import importlib.util
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from .storage import MaterialError


def soffice_path() -> str:
    configured = os.environ.get("CLONOTH_MATERIALS_SOFFICE", "").strip()
    if configured and Path(configured).is_file():
        return configured
    return shutil.which("soffice") or shutil.which("soffice.com") or ""


def capabilities() -> dict[str, Any]:
    modules = {name: importlib.util.find_spec(module) is not None for name, module in (
        ("pdf", "pypdf"), ("docx", "docx"), ("xlsx", "openpyxl"),
        ("pptx", "pptx"), ("pdf_build", "reportlab"), ("pdf_preview", "pypdfium2"),
    )}
    return {"modules": modules, "office_renderer": bool(soffice_path()),
            "ocr": bool(os.environ.get("CLONOTH_MATERIALS_TESSERACT") or shutil.which("tesseract")),
            "max_file_bytes": 50 * 1024 * 1024, "max_pages": 200,
            "max_export_pages": 50, "preview_confirmation_required": True}


def office_convert(path: Path, output: Path, target_format: str = "pdf") -> Path:
    binary = soffice_path()
    if not binary:
        raise MaterialError("缺少 Office 渲染器。请配置 CLONOTH_MATERIALS_SOFFICE 后重试；成品尚未发送。", "renderer_unavailable", 503)
    output.mkdir(parents=True, exist_ok=True)
    # Office filters still hit Windows path limits inside deeply nested job folders.
    with tempfile.TemporaryDirectory(prefix="co-", ignore_cleanup_errors=True) as temporary:
        temporary_root = Path(temporary)
        profile = temporary_root / "profile"
        user = profile / "user"
        user.mkdir(parents=True)
        incoming = temporary_root / ("input" + path.suffix)
        shutil.copyfile(path, incoming)
        converted = temporary_root / "converted"
        converted.mkdir()
        (user / "registrymodifications.xcu").write_text(
            '<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item>'
            '</oor:items>', encoding="utf-8",
        )
        try:
            result = subprocess.run(
                [binary, "--headless", "--nologo", "--nodefault", "--norestore",
                 f"-env:UserInstallation={profile.as_uri()}", "--convert-to", target_format, "--outdir", str(converted), str(incoming)],
                cwd=temporary_root, capture_output=True, timeout=120, check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise MaterialError("Office 渲染超时，请减少文档大小后重试。", "render_timeout", 504) from exc
        extension = target_format.split(":", 1)[0]
        produced = converted / ("input." + extension)
        if result.returncode or not produced.is_file() or not produced.stat().st_size:
            raise MaterialError("Office 渲染失败，成品尚未发送。", "render_failed", 422)
        pdf = output / (path.stem + "." + extension)
        shutil.copyfile(produced, pdf)
        return pdf


def recalculate_workbook(path: Path, directory: Path) -> Path:
    from openpyxl import load_workbook

    calculated = office_convert(path, directory, "xlsx:Calc MS Excel 2007 XML")
    formulas = load_workbook(calculated, data_only=False, keep_links=False)
    values = load_workbook(calculated, data_only=True, keep_links=False)
    try:
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    result = values[sheet.title][cell.coordinate]
                    if result.data_type == "e" or result.value is None:
                        raise MaterialError(f"公式 {sheet.title}!{cell.coordinate} 未得到有效计算结果。", "formula_error", 422)
    finally:
        formulas.close()
        values.close()
    return calculated


def render_preview(path: Path, directory: Path) -> dict:
    from PIL import Image

    directory.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
        with Image.open(path) as picture:
            if picture.width * picture.height > 40_000_000:
                raise MaterialError("图片像素超过 4000 万。", "limit_exceeded", 413)
            image = picture.convert("RGB")
            image.thumbnail((1600, 1600))
            target = directory / "page-1.png"
            image.save(target)
        return {"pages": [str(target)], "page_count": 1, "format": "image"}
    pdf = path if path.suffix.lower() == ".pdf" else office_convert(path, directory)
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise MaterialError("缺少 PDF 预览依赖，请安装 requirements-materials.txt。", "renderer_unavailable", 503) from exc
    pages = []
    document = pdfium.PdfDocument(pdf)
    try:
        if len(document) > 50:
            raise MaterialError("成品超过 50 页预览上限，请减少内容后重试。", "limit_exceeded", 413)
        for index in range(len(document)):
            page = document[index]
            try:
                width, height = page.get_size()
                scale = min(1.5, 1800 / max(width, height))
                bitmap = page.render(scale=scale)
                try:
                    picture = bitmap.to_pil()
                    target = directory / f"page-{index + 1}.png"
                    picture.save(target)
                    pages.append(str(target))
                finally:
                    bitmap.close()
            finally:
                page.close()
    finally:
        document.close()
    if not pages:
        raise MaterialError("成品没有可预览页面。", "render_failed", 422)
    return {"pages": pages, "page_count": len(pages), "format": "pdf", "pdf": str(pdf)}
