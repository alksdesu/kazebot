from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path
from typing import Any

from .storage import MAX_TEXT_CHARS, MaterialError

MAX_PAGES = 200


def inspect_package(path: Path) -> None:
    if path.suffix.lower() not in {".docx", ".xlsx", ".pptx"}:
        return
    with zipfile.ZipFile(path) as package:
        entries = package.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 200 * 1024 * 1024:
            raise MaterialError("Office 压缩包展开后过大。", "limit_exceeded", 413)
        for entry in entries:
            name = entry.filename.lower()
            if "vbaproject" in name or "/embeddings/" in name:
                raise MaterialError("不接受包含宏或嵌入对象的文档。", "active_content")
            if entry.file_size > 1024 * 1024 and entry.file_size > max(entry.compress_size, 1) * 200:
                raise MaterialError("Office 文件压缩比例超出限制。", "limit_exceeded", 413)
            if name.endswith(".rels"):
                text = package.read(entry).decode("utf-8", "replace")
                if 'TargetMode="External"' in text and any(token in text for token in ("/image", "/attachedTemplate", "/oleObject", "/externalLink")):
                    raise MaterialError("文档含外部资源，请先移除外链资源。", "active_content")


def parse_document(path: Path) -> dict[str, Any]:
    inspect_package(path)
    spans: list[dict] = []
    total_chars = 0

    def add(text: str, locator: dict) -> None:
        nonlocal total_chars
        if not text.strip():
            return
        total_chars += len(text)
        if total_chars > MAX_TEXT_CHARS or len(spans) >= 20000:
            raise MaterialError("文档文本超出解析上限，请拆分后上传。", "limit_exceeded", 413)
        spans.append({"index": len(spans) + 1, "text": text, "locator": locator})

    ext = path.suffix.lower()
    info: dict[str, Any] = {"spans": spans, "warnings": []}
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise MaterialError("PDF 已加密，请上传已解锁的副本。", "encrypted")
        if len(reader.pages) > MAX_PAGES:
            raise MaterialError(f"PDF 超过 {MAX_PAGES} 页，请拆分后上传。", "limit_exceeded", 413)
        empty_pages = []
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text(extraction_mode="layout") or ""
            add(text, {"kind": "pdf_page", "page": number, "label": f"PDF 第 {number} 页"})
            if not text.strip():
                empty_pages.append(number)
        info["pages"] = len(reader.pages)
        if empty_pages:
            info["warnings"].append({"code": "ocr_required", "pages": empty_pages, "message": "这些页面未提取到文字，需 OCR 或视觉识别；不能视为内容为空。"})
    elif ext == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(path)
        paragraph = table = 0
        for block in document.element.body:
            if block.tag.endswith("}p"):
                paragraph += 1
                add(Paragraph(block, document).text, {"kind": "docx_paragraph", "paragraph": paragraph, "label": f"正文第 {paragraph} 段"})
            elif block.tag.endswith("}tbl"):
                table += 1
                for row_i, row in enumerate(Table(block, document).rows, 1):
                    for col_i, cell in enumerate(row.cells, 1):
                        add(cell.text, {"kind": "docx_cell", "table": table, "row": row_i, "column": col_i, "label": f"表 {table} 第 {row_i} 行第 {col_i} 列"})
        info["pagination"] = "paragraphs"
        info["warnings"].append({"code": "no_word_page_numbers", "message": "Word 原件页码取决于排版环境，此处引用真实段落/表格位置。"})
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        count = 0
        try:
            if len(workbook.sheetnames) > 30:
                raise MaterialError("工作表超过 30 张。", "limit_exceeded", 413)
            for sheet in workbook:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        count += 1
                        if count > 50000:
                            raise MaterialError("有效单元格超过 50000 个。", "limit_exceeded", 413)
                        add(str(cell.value), {"kind": "sheet_cell", "sheet": sheet.title, "cell": cell.coordinate, "label": f"{sheet.title}!{cell.coordinate}"})
        finally:
            workbook.close()
    elif ext == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        if len(presentation.slides) > MAX_PAGES:
            raise MaterialError("幻灯片超过 200 页。", "limit_exceeded", 413)
        for number, slide in enumerate(presentation.slides, 1):
            for shape_i, shape in enumerate(slide.shapes, 1):
                if shape.has_text_frame:
                    add(shape.text, {"kind": "slide", "slide": number, "shape": shape_i, "label": f"幻灯片 {number}，文本框 {shape_i}"})
                if shape.has_table:
                    for row_i, row in enumerate(shape.table.rows, 1):
                        for col_i, cell in enumerate(row.cells, 1):
                            add(cell.text, {"kind": "slide_cell", "slide": number, "row": row_i, "column": col_i, "label": f"幻灯片 {number} 表格 {row_i}/{col_i}"})
    elif ext in {".txt", ".md", ".csv", ".tsv", ".json"}:
        text = path.read_text(encoding="utf-8-sig")
        if ext in {".csv", ".tsv"}:
            for number, row in enumerate(csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else ","), 1):
                add("\t".join(row), {"kind": "text_row", "row": number, "label": f"第 {number} 行"})
        else:
            for number, line in enumerate(text.splitlines(), 1):
                add(line, {"kind": "text_line", "line": number, "label": f"第 {number} 行"})
    else:
        raise MaterialError("该文件类型尚不支持文字解析，请使用 PDF、DOCX、XLSX、PPTX 或 UTF-8 文本。", "unsupported_type", 415)
    return info
