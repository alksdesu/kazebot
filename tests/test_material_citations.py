from __future__ import annotations

from pathlib import Path

import pytest

from engine.materials import MaterialError, MaterialService
from supervisor.feature_auth import FeatureActor


def test_pdf_page_and_word_paragraph_are_real_reopenable_citations(tmp_path: Path):
    from docx import Document
    from reportlab.pdfgen.canvas import Canvas

    uploads = tmp_path / "data/attachments"
    uploads.mkdir(parents=True)
    pdf = uploads / "reference.pdf"
    canvas = Canvas(str(pdf))
    canvas.drawString(50, 750, "First page is only background.")
    canvas.showPage()
    canvas.drawString(50, 750, "Deployment requires version 42.")
    canvas.showPage()
    canvas.save()
    word = uploads / "reference.docx"
    document = Document()
    document.add_paragraph("First paragraph is background.")
    document.add_paragraph("Deployment requires version 43.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Checksum"
    table.cell(0, 1).text = "fixture-only"
    document.save(word)
    actor = FeatureActor(scope="qq_group:one", owner="tester", bot_scope="bot")
    service = MaterialService(tmp_path)
    try:
        pdf_source = service.register_source(actor, {"path": "data/attachments/reference.pdf"})
        pdf_hit = service.search_source(actor, pdf_source["id"], "version 42")["results"][0]
        assert pdf_hit["locator"] == {"kind": "pdf_page", "page": 2, "label": "PDF 第 2 页"}
        quote = service.cite(actor, pdf_source["id"], pdf_hit["id"], "Deployment requires version 42.")
        assert quote["quote"] == "Deployment requires version 42."
        word_source = service.register_source(actor, {"path": "data/attachments/reference.docx"})
        word_hit = service.search_source(actor, word_source["id"], "version 43")["results"][0]
        assert word_hit["locator"]["paragraph"] == 2
        table_hit = service.search_source(actor, word_source["id"], "fixture-only")["results"][0]
        assert table_hit["locator"] == {"kind": "docx_cell", "table": 1, "row": 1, "column": 2, "label": "表 1 第 1 行第 2 列"}
        with pytest.raises(MaterialError):
            service.cite(actor, pdf_source["id"], pdf_hit["id"], "Deployment requires version 99.")
        source_id, span_id = pdf_source["id"], pdf_hit["id"]
    finally:
        service.store.close()
    reopened = MaterialService(tmp_path)
    try:
        assert reopened.cite(actor, source_id, span_id, "version 42")["quote"] == "version 42"
    finally:
        reopened.store.close()
