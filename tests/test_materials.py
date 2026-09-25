from __future__ import annotations

import asyncio
import copy
import importlib.util
import json
import os
import threading
import time
from collections import OrderedDict
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from fastapi import FastAPI

from engine.materials import MaterialError, MaterialService
from engine.materials.integration import begin_image_generation, persist_generated_images
from supervisor.feature_auth import FeatureActor
from supervisor.materials_api import MaterialRuntime, create_router
import supervisor.admin_api as admin_api


@pytest.fixture()
def actor():
    return FeatureActor(scope="qq_group:test", owner="user:test", is_admin=True, role="admin", bot_scope="bot", channel="qq_group", message_id="101")


@pytest.fixture()
def service(tmp_path):
    instance = MaterialService(tmp_path)
    yield instance
    instance.store.close()


def source(service, actor, name="sample.txt", content=b"first line\nsecond exact line"):
    path = service.workspace / "data" / "attachments" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    return service.register_source(actor, {"path": path.relative_to(service.workspace).as_posix(), "name": name})


def image_source(service, actor, name="sample.png", color="red"):
    from PIL import Image
    path = service.workspace / "data" / "attachments" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (64, 48), color).save(path)
    return service.register_source(actor, {"path": path.relative_to(service.workspace).as_posix(), "name": name})


def test_source_is_persistent_scoped_and_citations_are_exact(service, actor):
    item = source(service, actor)
    (service.workspace / "data" / "attachments" / "sample.txt").unlink()
    found = service.search_source(actor, item["id"], "exact")
    assert found["results"][0]["locator"]["line"] == 2
    span = found["results"][0]
    assert service.cite(actor, item["id"], span["id"], "second exact line")["quote"] == "second exact line"
    with pytest.raises(MaterialError, match="不匹配"):
        service.cite(actor, item["id"], span["id"], "invented line")
    with pytest.raises(MaterialError) as failure:
        service.search_source(replace(actor, scope="qq_private:other"), item["id"])
    assert failure.value.status == 404


@pytest.mark.parametrize("format", ["pdf", "docx"])
def test_pdf_pages_and_word_paragraph_citations_use_actual_source_positions(service, actor, format):
    if importlib.util.find_spec("pypdf") is None or importlib.util.find_spec("docx") is None:
        pytest.skip("Install requirements-materials.txt for document parsing acceptance")
    path = service.workspace / "data" / "attachments" / f"evidence.{format}"
    path.parent.mkdir(parents=True, exist_ok=True)
    if format == "pdf":
        from reportlab.pdfgen.canvas import Canvas
        canvas = Canvas(str(path))
        canvas.drawString(72, 720, "first evidence")
        canvas.showPage()
        canvas.drawString(72, 720, "second evidence")
        canvas.save()
    else:
        from docx import Document
        document = Document()
        document.add_paragraph("first evidence")
        document.add_paragraph("second evidence")
        document.save(path)
    registered = service.register_source(actor, {"path": path.relative_to(service.workspace).as_posix()})
    span = service.search_source(actor, registered["id"], "second evidence")["results"][0]
    assert span["locator"]["page" if format == "pdf" else "paragraph"] == 2
    if format == "docx":
        assert "page" not in span["locator"]
    assert service.cite(actor, registered["id"], span["id"], "second evidence")["source_sha256"] == registered["sha256"]


def test_registration_rejects_workspace_files_and_path_escape(service, actor):
    private = service.workspace / "secret.txt"
    private.write_text("secret", encoding="utf-8")
    with pytest.raises(MaterialError) as failure:
        service.register_source(actor, {"path": str(private)})
    assert failure.value.status == 403


def test_journal_opt_in_scope_retention_original_text_and_attachment_patch(service, actor):
    body = {"message_id": "101", "text": "exact original", "timestamp": time.time(), "sender_name": "Display"}
    assert service.record_message(actor, body)["recorded"] is False
    service.store.save_settings(actor, {"journal_enabled": True})
    first = service.record_message(actor, body)
    attachment = source(service, actor)
    again = service.record_message(actor, {**body, "text": "do not rewrite", "attachments": [{"source_id": attachment["id"]}]})
    assert again["id"] == first["id"]
    assert again["text"] == "exact original"
    assert again["attachments"][0]["source_id"] == attachment["id"]
    assert service.search_messages(actor, "original")["results"][0]["message_id"] == "101"
    assert service.get_message(actor, "101")["timestamp"] == body["timestamp"]
    assert service.search_messages(replace(actor, scope="other"), "original")["results"] == []
    service.store.save_settings(actor, {"journal_enabled": False})
    with pytest.raises(MaterialError):
        service.get_message(actor, "101")


def test_journal_preserves_original_whitespace_and_source_origins(service, actor):
    service.store.save_settings(actor, {"journal_enabled": True})
    original = "  indented original\n\ttrailing whitespace  \n"
    service.record_message(actor, {"message_id": "101", "text": original, "timestamp": time.time()})
    assert service.get_message(actor, "101")["text"] == original
    first = source(service, actor)
    second = source(service, replace(actor, owner="other-user", message_id="102"))
    assert first["id"] != second["id"]
    assert first["sha256"] == second["sha256"]


def test_extraction_preserves_zero_and_requires_matching_revision(service, actor):
    image = image_source(service, actor)
    extracted = service.create_extraction(actor, {"source_id": image["id"], "kind": "table", "fields": [
        {"id": "price", "label": "金额", "value": 0, "row": 2, "column": 1, "bbox": [.1, .1, .9, .9]},
    ]})
    assert extracted["fields"][0]["value"] == "0"
    fixed = service.correct_extraction(actor, extracted["id"], {"expected_revision": 1, "values": {"price": "120.50"}, "confirm": True})
    assert fixed["status"] == "reviewed"
    assert fixed["revisions"][0]["fields"][0]["value"] == "0"
    with pytest.raises(MaterialError) as conflict:
        service.correct_extraction(actor, extracted["id"], {"expected_revision": 1, "values": {"price": "3"}})
    assert conflict.value.code == "version_conflict"
    with pytest.raises(MaterialError):
        service.create_extraction(actor, {"source_id": image["id"], "fields": [{"label": "bad", "value": "x", "bbox": [0, 0, 2, 1]}]})


def test_versions_restore_original_bytes_and_block_unconfirmed_exports(service, actor):
    one = image_source(service, actor, color="red")
    v1 = service.register_image_version(actor, {"source_id": one["id"]})
    two = image_source(service, actor, name="second.png", color="blue")
    v2 = service.register_image_version(actor, {"source_id": two["id"], "artifact_id": v1["artifact_id"], "expected_version": v1["id"]})
    with pytest.raises(MaterialError) as blocked:
        service.export_version(actor, v2["id"])
    assert blocked.value.code == "approval_required"
    restored = service.restore_version(actor, v1["artifact_id"], {"version_id": v1["id"], "expected_version": v2["id"]})
    assert restored["sha256"] == v1["sha256"]
    assert restored["number"] == 3
    with pytest.raises(MaterialError):
        service.approve_version(actor, v2["id"], {"sha256": v2["sha256"], "preview_inspected": True})
    service.approve_version(actor, restored["id"], {"sha256": restored["sha256"], "preview_inspected": True})
    exported = service.export_version(actor, restored["id"])
    assert (service.workspace / exported["attachment"]["path"]).read_bytes() == (service.workspace / "data" / "attachments" / "sample.png").read_bytes()


def generated_body(service, color="red", filename="generated.png"):
    from PIL import Image
    path = service.workspace / "data" / "attachments" / "gpt_image" / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (64, 48), color).save(path)
    return {"tool_name": "gpt_image_2", "tool_call_id": "call1", "attachments": [{"type": "image", "path": path.relative_to(service.workspace).as_posix(), "name": filename}]}


def test_image_operation_completion_is_bound_idempotent_and_chains(service, actor):
    actor = replace(actor, task_id="origin", interactive=False)
    body = generated_body(service)
    operation = service.begin_image_operation(actor, body)
    v1 = service.complete_image_operation(operation["id"], "origin", body)["versions"][0]
    assert service.complete_image_operation(operation["id"], "origin", body)["versions"][0]["id"] == v1["id"]
    with pytest.raises(MaterialError):
        service.complete_image_operation(operation["id"], "other-task", body)
    original = service.store.get(v1["id"], actor, kind="version")
    next_body = {**generated_body(service, "blue", "second.png"), "tool_call_id": "call2", "image_paths": [str(service.store.path(original["blob"]))]}
    second_operation = service.begin_image_operation(actor, next_body)
    v2 = service.complete_image_operation(second_operation["id"], "origin", next_body)["versions"][0]
    assert v2["artifact_id"] == v1["artifact_id"]
    assert v2["metadata"]["base_version"] == v1["id"]
    assert v2["number"] == 2


def test_document_previews_cannot_be_archived_as_generated_art(service, actor):
    actor = replace(actor, task_id="origin", interactive=False)
    image = image_source(service, actor)
    operation = service.begin_image_operation(actor, {"tool_name": "gpt_image_2"})
    with pytest.raises(MaterialError) as denied:
        service.complete_image_operation(operation["id"], "origin", {"attachments": [{"path": "data/attachments/sample.png"}]})
    assert denied.value.status == 403


def test_integration_reserves_once_per_call_and_preserves_real_result_on_archive_failure():
    calls = []

    async def handler(request):
        calls.append(str(request.url))
        if request.url.path.endswith("/image-operations"):
            return httpx.Response(200, json={"id": f"operation-{len(calls)}"})
        return httpx.Response(503, json={"detail": "archive unavailable"})

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            ctx = SimpleNamespace(http=http, task_id="t", tool_call_id="c1", supervisor_url="http://test")
            a = await begin_image_generation(ctx, "gpt_image_2", {"prompt": "test"})
            b = await begin_image_generation(ctx, "gpt_image_2", {"prompt": "test"})
            assert a == b and len(calls) == 1
            ctx.tool_call_id = "c2"
            c = await begin_image_generation(ctx, "gpt_image_2", {"prompt": "test"})
            assert c["id"] != a["id"]
            real = {"ok": True, "data": {"result": "generated", "attachments": [{"type": "image", "path": "data/attachments/gpt_image/a.png"}]}}
            result = await persist_generated_images(ctx, "gpt_image_2", {}, real, operation=c)
            assert result["ok"] is True
            assert result["data"]["attachments"]
            assert result["data"]["materials_status"] == "archive_failed"
            assert "不要再次付费" in result["data"]["result"]

    asyncio.run(exercise())


@pytest.fixture()
def api_state(tmp_path, monkeypatch):
    monkeypatch.setattr(admin_api, "_admin_token", "materials-test")
    monkeypatch.setattr(admin_api, "_auth_failures", OrderedDict())
    task = SimpleNamespace(cancel_requested=False, status="running", input={"attachments": [], "task_context": {
        "conversation_key": "qq_group:test", "channel": "qq_group", "platform_auth": {"platform": "qq", "user_id": "123", "bot_scope": "bot"},
    }})
    state = SimpleNamespace(workspace_root=tmp_path, _lock=threading.RLock(), tasks={"task": task}, _task_terminal=lambda item: item.status == "finished")
    app = FastAPI()
    app.include_router(create_router(state))
    yield app, state
    state.materials.store.close()


def test_model_cannot_confirm_or_forge_platform_journal(api_state):
    app, _ = api_state

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test", headers={"Authorization": "Bearer materials-test", "X-Clonoth-Task-Id": "task"}) as http:
            assert (await http.post("/v1/materials/versions/v_fake/approve", json={"sha256": "anything", "preview_inspected": True})).status_code == 403
            assert (await http.post("/v1/materials/journal/messages", json={"message_id": "42", "text": "fake original", "timestamp": time.time()})).status_code == 403
            assert (await http.post("/v1/materials/sources/register", json={"path": "data/attachments/other.png"})).status_code == 403

    asyncio.run(exercise())


def test_finished_origin_can_complete_only_its_reserved_image_operation(api_state):
    app, state = api_state
    body = generated_body(state.materials)

    async def exercise():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test", headers={"Authorization": "Bearer materials-test", "X-Clonoth-Task-Id": "task"}) as http:
            operation = await http.post("/v1/materials/image-operations", json=body)
            assert operation.status_code == 200
            state.tasks["task"].status = "finished"
            assert (await http.post("/v1/materials/image-operations", json=body)).status_code == 403
            completed = await http.post(f"/v1/materials/image-operations/{operation.json()['id']}/complete", json=body)
            assert completed.status_code == 200
            assert completed.json()["versions"][0]["metadata"]["tool"] == "gpt_image_2"
            state.tasks["task"].cancel_requested = True
            assert (await http.post(f"/v1/materials/image-operations/{operation.json()['id']}/complete", json=body)).status_code == 403

    asyncio.run(exercise())


@pytest.mark.parametrize("format", ["pdf", "docx", "xlsx", "pptx"])
def test_real_generation_render_confirmation_and_reopen(service, actor, format):
    if importlib.util.find_spec("pypdfium2") is None:
        pytest.skip("Install requirements-materials.txt for renderer acceptance")
    if format != "pdf" and not os.environ.get("CLONOTH_MATERIALS_SOFFICE"):
        pytest.skip("Set CLONOTH_MATERIALS_SOFFICE for real Office renderer acceptance")
    spec = {"title": "Synthetic acceptance", "sections": [{"heading": "Evidence", "paragraphs": ["Synthetic document only."], "table": [["Item", "Value"], ["Total", "6"]]}]}
    if format == "xlsx":
        spec = {"title": "Synthetic", "sheets": [{"name": "Totals", "rows": [["A", "B", "Product"], [2, 3, {"formula": "=A2*B2"}]]}]}
    if format == "pptx":
        spec = {"title": "Synthetic", "slides": [{"title": "Synthetic acceptance", "bullets": ["Editable text", "Preview before send"]}]}
    job = service.run_job(actor, service.create_job(actor, {"format": format, "spec": spec})["id"])
    assert job["status"] == "preview_ready", job.get("error")
    version = service.store.get(job["version_id"], actor, kind="version")
    from PIL import Image
    with Image.open(service.preview_file(actor, version["id"], 1)) as image:
        assert image.width > 100 and image.height > 100
    with pytest.raises(MaterialError):
        service.export_version(actor, version["id"])
    service.approve_version(actor, version["id"], {"sha256": version["sha256"], "preview_inspected": True})
    exported = service.export_version(actor, version["id"])
    path = service.workspace / exported["attachment"]["path"]
    assert path.is_file()
    if format == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, data_only=True)
        assert workbook["Totals"]["C2"].value == 6
        workbook.close()


def test_missing_renderer_never_marks_office_output_preview_ready(service, actor, monkeypatch):
    import engine.materials.service as module

    def unavailable(*args, **kwargs):
        raise MaterialError("renderer unavailable", "renderer_unavailable", 503)

    monkeypatch.setattr(module, "run_worker", unavailable)
    job = service.run_job(actor, service.create_job(actor, {"format": "docx", "spec": {"title": "Test", "sections": [{"paragraphs": ["test"]}]}})["id"])
    assert job["status"] == "failed"
    assert job["error_code"] == "renderer_unavailable"
    assert "version_id" not in job


def test_restart_marks_unfinished_local_jobs_interrupted_and_retry_is_explicit(tmp_path, actor):
    first = MaterialService(tmp_path)
    spec = {"title": "Restart", "sections": [{"paragraphs": ["Synthetic"]}]}
    queued = first.create_job(actor, {"format": "pdf", "spec": spec})
    running = first.create_job(actor, {"format": "docx", "spec": spec})
    running["status"] = "running"
    first.store.update(running)
    first.close()
    second = MaterialService(tmp_path)
    try:
        assert second.store.get(queued["id"], actor)["status"] == "interrupted"
        assert second.store.get(running["id"], actor)["status"] == "interrupted"
        retry = second.retry_job(actor, running["id"])
        assert retry["id"] != running["id"]
        assert retry["status"] == "queued"
        assert retry["artifact_id"] == running["artifact_id"]
        assert second.store.list("version", actor) == []
    finally:
        second.close()


def test_restart_recovers_already_committed_preview_without_regenerating(tmp_path, actor):
    first = MaterialService(tmp_path)
    job = first.create_job(actor, {"format": "pdf", "spec": {"title": "Test", "sections": []}})
    image = image_source(first, actor)
    source_record = first.source(actor, image["id"])
    path = first.store.path(source_record["blob"])
    saved = first._save_version(actor, job["artifact_id"], path, {"pages": [str(path)]}, expected="", metadata={"job_id": job["id"]}, job_id=job["id"])
    job["status"] = "running"
    first.store.update(job)
    first.close()
    second = MaterialService(tmp_path)
    try:
        recovered = second.store.get(job["id"], actor, kind="job")
        assert recovered["status"] == "preview_ready"
        assert recovered["version_id"] == saved["id"]
        assert len(second.store.list("version", actor)) == 1
        assert second.store.get(saved["id"], actor)["approved_hash"] == ""
    finally:
        second.close()


def test_shutdown_waits_for_worker_completion_before_closing_database(service, actor, monkeypatch):
    entered = threading.Event()
    release = threading.Event()

    def blocked_worker(*args, **kwargs):
        entered.set()
        assert release.wait(3)
        raise MaterialError("worker stopped", "interrupted", 409)

    monkeypatch.setattr(service, "_work", blocked_worker)
    job = service.create_job(actor, {"format": "pdf", "spec": {"title": "Close", "sections": []}})
    runtime = MaterialRuntime(service)

    async def exercise():
        runtime.launch(actor, job)
        assert await asyncio.to_thread(entered.wait, 2)
        closing = asyncio.create_task(runtime.close())
        await asyncio.sleep(.03)
        assert not closing.done()
        assert service.store.get(job["id"], actor)["status"] == "interrupted"
        release.set()
        await closing
        assert runtime.closed
        await runtime.close()

    asyncio.run(exercise())
    reopened = MaterialService(service.workspace)
    try:
        assert reopened.store.get(job["id"], actor)["status"] == "interrupted"
        assert reopened.store.list("version", actor) == []
    finally:
        reopened.close()
